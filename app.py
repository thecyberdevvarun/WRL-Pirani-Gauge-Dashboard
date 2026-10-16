from flask import (
    Flask,
    request,
    jsonify,
    send_file,
    send_from_directory,
    Response,
    session,
    g,
)
from test_runner import get_active_tests, run_test, stop_test, MAX_DURATION_MIN
from db import (
    get_connection,
    ensure_test_header_columns,
    get_material_name,
)
from lines_config import ConfigError, load_config_from_environment, warnings_for
from polling import PollingService, SingleOwnerLease
from register_maps import REGISTER_MAPS, SUCCESS, read_gauge
import pandas as pd
import io
import time
import threading
import queue
import logging
import os
import atexit
import hmac
import secrets
from datetime import datetime, timedelta
from functools import wraps
from cachetools import TTLCache
import json
from pymodbus.client.sync import ModbusTcpClient
import sys
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from openpyxl.drawing.image import Image as XLImage
import openpyxl
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.shapes import Drawing, String

logging.basicConfig(
    level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s"
)
logger = logging.getLogger(__name__)

# The React app (client/) is built to client/dist and served as static
# files by this same Flask process in production. See client/README.md.
CLIENT_DIST = os.path.join(os.path.dirname(os.path.abspath(__file__)), "client", "dist")

app = Flask(__name__, static_folder=None)

# =========================================================
# PRODUCTION LINES — everything comes from .env (see lines_config.py)
# =========================================================
# Lines (label, gauge count, USRs, upper limit, poll interval, IN/OUT gap,
# network), users and polling settings. Loaded once at startup; edit .env and
# restart the app to change them. A bad .env stops the app from starting and
# prints every problem found.
try:
    CONFIG = load_config_from_environment()
except ConfigError as exc:
    logger.critical("%s", exc)
    raise

LINES = CONFIG.lines         # LINE_KEY -> lines_config.Line
USERS = CONFIG.users         # lower-case username -> lines_config.User
SETTINGS = CONFIG.settings   # polling mode, timeouts, backoff, ...

# =========================================================
# SESSIONS — the backend, not the browser, decides which line you are on
# =========================================================
SECRET_KEY = os.environ.get("SECRET_KEY", "").strip()
if not SECRET_KEY:
    SECRET_KEY = secrets.token_hex(32)
    logger.warning(
        "SECRET_KEY is not set in .env — using a temporary key, so everyone is "
        "logged out whenever the app restarts."
    )
try:
    SESSION_HOURS = float(os.environ.get("SESSION_HOURS", "12"))
except ValueError:
    SESSION_HOURS = 12.0

app.secret_key = SECRET_KEY
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=SESSION_HOURS),
)

# =========================================================
# GLOBALS
# =========================================================
# Every gauge is identified by (line_key, slave_id): the slave ID is the gauge
# number *within* its line, so "gauge 1" exists on every line.
#
# POLLING:       managed Modbus polling (one worker per USR) + live gauge status
# FIXTURE_CACHE: line_key -> [fixture state dicts, one per enabled gauge]
# POLL_LEASE:    only one process on this PC may poll the gauges
FIXTURE_CACHE = {}
POLLING = PollingService(LINES, SETTINGS)
POLL_LEASE = SingleOwnerLease(SETTINGS.lease_port)
REPORT_CACHE = TTLCache(maxsize=200, ttl=30)
LOCK = threading.Lock()

# Company logo, embedded in the app header/login page (via the built
# frontend) and in PDF/Excel report exports (directly, below).
LOGO_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "logo", "Western-Refrigeration-768x453.jpg"
)

# Per-client SSE queues: (line_key, queue) — one per connected browser tab,
# so a client only ever receives the fixtures of the line it is logged into.
SSE_CLIENTS = []
SSE_LOCK = threading.Lock()


# =========================================================
# AUTH
# =========================================================
# A login is a User: one or more authorised lines and a role. Every line has
# its own login (one line -> selected automatically); extra USERS may be
# authorised for several lines and then choose which one to work on. The
# chosen line lives in the signed session cookie and is re-checked against the
# user's authorised lines on every request.
ROLE_LABEL = {"operator": "Operator", "admin": "Admin"}


def current_user():
    """The logged-in User, or None (also None if removed from .env since login)."""
    return USERS.get((session.get("username") or "").lower())


def current_line():
    """The Line this session is working on, or None."""
    user = current_user()
    key = session.get("line_key")
    if user is None or key is None or key not in user.lines:
        return None
    return LINES.get(key)


def user_required(fn):
    """Logged in; a production line does not have to be selected yet."""

    @wraps(fn)
    def wrapper(*args, **kwargs):
        user = current_user()
        if user is None:
            session.clear()
            return jsonify({"error": "Not logged in"}), 401
        g.user = user
        return fn(*args, **kwargs)

    return wrapper


def login_required(fn):
    """Logged in with a line selected; exposes g.user and g.line.

    Every data endpoint is scoped to g.line — the line is never taken from the
    request, so a user cannot read or control a line they are not authorised for.
    """

    @wraps(fn)
    def wrapper(*args, **kwargs):
        user = current_user()
        if user is None:
            session.clear()
            return jsonify({"error": "Not logged in"}), 401
        line = current_line()
        if line is None:
            return (
                jsonify({"error": "No production line selected", "code": "line_not_selected"}),
                409,
            )
        g.user, g.line = user, line
        return fn(*args, **kwargs)

    return wrapper


def admin_required(fn):
    """Use below @login_required: only users with the admin role may continue."""

    @wraps(fn)
    def wrapper(*args, **kwargs):
        if g.user.role != "admin":
            return jsonify({"error": "This action needs an admin login"}), 403
        return fn(*args, **kwargs)

    return wrapper


def _match_user(username, password):
    """Check credentials against every user, without exiting early on a match."""
    u = username.strip().lower().encode()
    p = password.encode()
    match = None
    for user in USERS.values():
        user_ok = hmac.compare_digest(u, user.username.lower().encode())
        pass_ok = hmac.compare_digest(p, user.password.encode())
        if user_ok and pass_ok:
            match = user
    return match


def _session_payload(user, line):
    return {
        "username": user.username,
        "role": ROLE_LABEL[user.role],
        "line": line.key if line else None,
        "line_label": line.label if line else None,
        "gauge_count": len(line.gauge_ids) if line else None,
        "lines": [
            {"key": k, "label": LINES[k].label, "gauge_count": len(LINES[k].gauge_ids)}
            for k in user.lines
            if k in LINES
        ],
        "needs_line_selection": line is None,
    }


@app.route("/api/auth/login", methods=["POST"])
def auth_login():
    d = request.get_json(silent=True) or {}
    username = d.get("username")
    password = d.get("password")
    if not isinstance(username, str) or not isinstance(password, str):
        return jsonify({"error": "Username and password are required"}), 400

    user = _match_user(username, password)
    if user is None:
        logger.warning("Failed login attempt for username '%s'", username.strip()[:50])
        return jsonify({"error": "Invalid username or password"}), 401

    session.clear()
    session["username"] = user.username
    if len(user.lines) == 1:
        session["line_key"] = user.lines[0]          # single line: no chooser needed
    session.permanent = True
    logger.info("Login: user=%s lines=%s", user.username, ",".join(user.lines))
    return jsonify(_session_payload(user, current_line()))


@app.route("/api/auth/logout", methods=["POST"])
def auth_logout():
    # Ends this browser session only. Polling workers belong to the backend and
    # keep running for everyone else.
    session.clear()
    return jsonify({"success": True})


@app.route("/api/auth/me")
@user_required
def auth_me():
    return jsonify(_session_payload(g.user, current_line()))


@app.route("/api/auth/select-line", methods=["POST"])
@user_required
def auth_select_line():
    d = request.get_json(silent=True) or {}
    key = d.get("line")
    key = key.strip().upper() if isinstance(key, str) else ""
    if key not in g.user.lines or key not in LINES:
        logger.warning("Line selection refused: user=%s requested=%s", g.user.username, key[:50])
        return jsonify({"error": "You are not authorised for that line"}), 403
    session["line_key"] = key
    logger.info("Line selected: user=%s line=%s", g.user.username, key)
    return jsonify(_session_payload(g.user, LINES[key]))


# =========================================================
# IN/OUT AUTO-STOP CYCLE
# =========================================================
def auto_stop_cycle_gauge(line, in_gauge_id):
    """After a new IN scan starts a test on in_gauge_id, this line's gap
    setting means every gauge in the `gap` positions immediately *ahead* of
    it is now in the exit window — scanning gauge 15 with gap=5 sweeps
    gauges 16, 17, 18, 19, 20 (wrapping around if needed). Any of those that
    still have a test running gets stopped, letting its naturally
    accumulated PASS/FAIL stand (not ABORTED) — same outcome as a normal
    timeout completion. Gauges in the window with no active test are simply
    a no-op. If a swept gauge's test hasn't yet reached its line's minimum
    test duration, stop_test() defers the actual stop rather than cutting it
    short — it still counts as "stopped" here (the request was accepted).
    Returns the list of gauge ids a stop was requested for (may be empty).
    Only ever touches gauges of `line` — other lines are never affected."""
    gap = line.gap
    order = line.gauge_ids

    if not gap or len(order) < 2 or in_gauge_id not in order:
        return []

    idx = order.index(in_gauge_id)
    n = len(order)
    span = min(gap, n - 1)  # never wrap far enough to include in_gauge_id itself

    stopped = []
    for offset in range(1, span + 1):
        target = order[(idx + offset) % n]
        if stop_test(line.key, target, manual=False):
            stopped.append(target)
            logger.info(
                "IN/OUT cycle: gauge %s scanned IN -> auto-stopping gauge %s "
                "(%s of %s positions ahead, gap=%s, line=%s)",
                in_gauge_id, target, offset, span, gap, line.key,
            )
    return stopped


# =========================================================
# DB AUTO-HEAL DECORATOR
# =========================================================
def db_safe(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        for _ in range(3):
            try:
                return fn(*args, **kwargs)
            except Exception as e:
                logger.warning("[DB RETRY] %s: %s", fn.__name__, e)
                time.sleep(0.5)
        return jsonify({"error": "Database unavailable"}), 500

    return wrapper


# =========================================================
# UI ROUTES — serve the built React app (client/dist) in production.
# In development the React app runs separately on Vite (npm run dev,
# localhost:5173) and proxies API calls here, so these routes are not
# used at all in dev. They only matter once `npm run build` has been
# run and client/dist exists.
# =========================================================
@app.route("/")
@app.route("/<path:path>")
def serve_react_app(path=""):
    if path and os.path.exists(os.path.join(CLIENT_DIST, path)):
        return send_from_directory(CLIENT_DIST, path)

    index_path = os.path.join(CLIENT_DIST, "index.html")
    if os.path.exists(index_path):
        return send_from_directory(CLIENT_DIST, "index.html")

    return (
        jsonify(
            {
                "error": "Frontend build not found",
                "hint": "Run `npm run build` inside client/, or use `npm run dev` "
                "in client/ for local development (http://localhost:5173).",
            }
        ),
        404,
    )


# =========================================================
# HEALTH CHECK
# =========================================================
@app.route("/api/health")
def health():
    body = {"status": "ok", "timestamp": datetime.now().isoformat()}

    # Gateway and worker details only for a logged-in session with a line, and
    # only for that line.
    line = current_line()
    if line is not None:
        usrs = POLLING.line_snapshot(line.key)
        gauges = [gs for u in usrs for gs in u["gauges"].values()]
        enabled = [u for u in usrs if u["enabled"]]
        body.update(
            {
                "gateways": [gw.address for gw in line.gateways],
                "gauges_responding": sum(1 for gs in gauges if gs["freshness"] == "valid"),
                "gauges_polled": sum(1 for gs in gauges if gs["communication_status"] is not None),
                "polling": {
                    "mode": SETTINGS.polling_mode,
                    "polled_here": POLLING.polled_here(line.key),
                    "lease_held": POLL_LEASE.held,
                    "usrs_enabled": len(enabled),
                    "usrs_connected": sum(1 for u in enabled if u["connection_status"] == "connected"),
                    "workers_alive": sum(1 for u in enabled if u["worker_alive"]),
                },
            }
        )
    return jsonify(body)


# =========================================================
# TODAY'S STATS
# =========================================================
@app.route("/api/stats/today")
@login_required
@db_safe
def today_stats():
    conn = get_connection()
    cur = conn.cursor()
    cur.execute(
        """
        SELECT
            COUNT(*)                                                        AS total,
            SUM(CASE WHEN final_result = 'PASS' THEN 1 ELSE 0 END)        AS pass_count,
            SUM(CASE WHEN final_result = 'FAIL' THEN 1 ELSE 0 END)        AS fail_count,
            SUM(CASE WHEN final_result IS NULL   THEN 1 ELSE 0 END)       AS running_count
        FROM pirani_test_header
        WHERE CAST(start_time AS DATE) = CAST(GETDATE() AS DATE)
          AND line_name = ?
        """,
        (g.line.key,),
    )
    row = cur.fetchone()
    conn.close()

    total, pass_c, fail_c, running = (
        row[0] or 0,
        row[1] or 0,
        row[2] or 0,
        row[3] or 0,
    )
    return jsonify(
        {
            "total": total,
            "pass": pass_c,
            "fail": fail_c,
            "running": running,
            "pass_rate": round(pass_c / total * 100, 1) if total > 0 else 0.0,
        }
    )


# =========================================================
# MATERIAL LOOKUP (external system, display-only — does not gate testing)
# =========================================================
@app.route("/api/material/<alt_name>")
@login_required
@db_safe
def material_lookup(alt_name):
    name = get_material_name(alt_name)
    return jsonify({"exists": name is not None, "name": name})


# =========================================================
# START TEST
# =========================================================
@app.route("/start-test", methods=["POST"])
@login_required
def start_test():
    line = g.line
    d = request.get_json(silent=True) or {}

    serial_no = (d.get("serial_no") or "").strip()
    model_code = (d.get("model_code") or "").strip()

    try:
        gauge_id = int(d.get("gauge_id", 0))
    except (ValueError, TypeError):
        return jsonify({"status": "ERROR", "message": "Invalid gauge ID"}), 400

    if not serial_no or not model_code:
        return (
            jsonify(
                {"status": "ERROR", "message": "Serial number and model code required"}
            ),
            400,
        )

    # The line comes from the login session, never from the request body.
    if gauge_id in line.disabled_gauges:
        return jsonify({"status": "ERROR", "message": f"Gauge {gauge_id} is disabled"}), 400
    gateway = line.gateway_for(gauge_id) if gauge_id in line.gauge_ids else None
    if gateway is None:
        return jsonify({"status": "ERROR", "message": f"Unknown Gauge ID {gauge_id}"}), 404

    # In line_specific mode a PC can only reach its own line's USRs.
    if not POLLING.polled_here(line.key):
        return (
            jsonify(
                {
                    "status": "ERROR",
                    "message": f"{line.label} is not controlled from this PC "
                    f"(polling mode: {SETTINGS.polling_mode}). Start tests from the "
                    f"{line.label} line PC.",
                }
            ),
            409,
        )

    # Display-only lookup in an external system — never blocks starting a test.
    try:
        model_name = get_material_name(model_code) or ""
    except Exception:
        logger.warning("Material lookup failed for model_code=%s", model_code, exc_info=True)
        model_name = ""

    result = run_test(
        serial_no=serial_no,
        model_code=model_code,
        model_name=model_name,
        line_name=line.key,
        slave_id=gauge_id,
        host=gateway.host,
        port=gateway.port,
        upper_limit=line.upper_limit,
        poll_interval_sec=line.poll_interval_sec,
        reading_delay_sec=line.reading_delay_sec,
        min_duration_sec=line.min_duration_sec,
        register_map=line.register_map,
    )

    if result.get("status") == "STARTED":
        auto_stop_cycle_gauge(line, gauge_id)

    return jsonify(result)


# =========================================================
# STOP TEST
# =========================================================
@app.route("/stop-test/<int:gauge_id>", methods=["POST"])
@login_required
def stop_test_route(gauge_id):
    if stop_test(g.line.key, gauge_id):
        return jsonify(
            {"status": "STOPPED", "message": f"Stop signal sent to Gauge {gauge_id}"}
        )
    return (
        jsonify(
            {"status": "NOT_RUNNING", "message": f"No active test on Gauge {gauge_id}"}
        ),
        404,
    )


@app.route("/api/active-tests")
@login_required
def active_tests_api():
    gauges = get_active_tests(g.line.key)
    return jsonify({"active_gauges": gauges, "count": len(gauges)})


# =========================================================
# FIXTURE BACKGROUND WORKER
# =========================================================
def _broadcast_fixtures(by_line):
    """Push each line's fixtures to the SSE clients logged into that line."""
    with SSE_LOCK:
        for line_key, q in SSE_CLIENTS:
            try:
                q.put_nowait(by_line.get(line_key, []))
            except queue.Full:
                pass  # Slow client — drop this update


def _fixture_state(gid, latest):
    """Status dict for one gauge, given the latest test row for it (or None)."""
    if latest is None:
        return {
            "slave_id": gid,
            "status": "IDLE",
            "color": "gray",
            "remaining": None,
            "line": None,
            "serial_no": None,
        }

    serial, result, line_name, remaining = latest

    if remaining is not None and remaining <= 0 and result is not None:
        status, color = ("IDLE", "gray")
        remaining = None
    elif result is None:
        status, color = ("RUNNING", "yellow")
        remaining = max(0, remaining or 0)
    elif result == "PASS":
        status, color = ("PASS", "green")
        remaining = None
    elif result == "FAIL":
        status, color = ("FAIL", "red")
        remaining = None
    else:
        status, color = ("ERROR", "orange")
        remaining = None

    return {
        "slave_id": gid,
        "status": status,
        "color": color,
        "remaining": remaining,
        "line": line_name,
        "serial_no": serial,
    }


def build_fixtures_by_line(rows):
    """rows: (line_name, gauge_id, serial_no, final_result, remaining_sec) — the
    latest test per (line, gauge). Returns {line_key: [fixture, ...]} with one
    entry per enabled gauge of every configured line; rows for a line that is
    not configured (e.g. old data) are ignored."""
    latest = {(r[0], r[1]): (r[2], r[3], r[0], r[4]) for r in rows}
    return {
        line.key: [
            _fixture_state(gid, latest.get((line.key, gid))) for gid in line.gauge_ids
        ]
        for line in LINES.values()
    }


def fixture_worker():
    conn = None
    cur = None

    line_keys = list(LINES.keys())
    placeholders = ",".join("?" * len(line_keys))
    # Latest test per (line, gauge): the same gauge number exists on every
    # line, so the line is part of the identity.
    query = f"""
        SELECT
            h.line_name,
            h.gauge_id,
            h.serial_no,
            h.final_result,
            DATEDIFF(
                SECOND,
                GETDATE(),
                DATEADD(MINUTE, ?, h.start_time)
            ) AS remaining_sec
        FROM pirani_test_header h
        WHERE h.line_name IN ({placeholders})
          AND h.start_time = (
            SELECT MAX(start_time)
            FROM pirani_test_header h2
            WHERE h2.gauge_id = h.gauge_id
              AND h2.line_name = h.line_name
        )
    """

    while True:
        try:
            if conn is None:
                conn = get_connection()
                cur = conn.cursor()

            cur.execute(query, (MAX_DURATION_MIN, *line_keys))
            by_line = build_fixtures_by_line(cur.fetchall())

            with LOCK:
                FIXTURE_CACHE.clear()
                FIXTURE_CACHE.update(by_line)

            _broadcast_fixtures(by_line)

        except Exception as e:
            logger.error("[FIXTURE WORKER] %s", e)
            try:
                if conn:
                    conn.close()
            except Exception:
                pass
            conn = None
            cur = None
            time.sleep(2)

        time.sleep(2)


# =========================================================
# FIXTURES REST API
# =========================================================
@app.route("/api/fixtures")
@login_required
def get_fixtures():
    with LOCK:
        data = list(FIXTURE_CACHE.get(g.line.key, []))
    return jsonify(data)


# =========================================================
# FIXTURES SSE (multi-client safe, one line per client)
# =========================================================
@app.route("/api/fixtures/stream")
@login_required
def fixtures_stream():
    line_key = g.line.key
    client_q = queue.Queue(maxsize=10)
    client = (line_key, client_q)
    with SSE_LOCK:
        SSE_CLIENTS.append(client)

    def stream():
        # Push current state immediately on connect
        with LOCK:
            snapshot = list(FIXTURE_CACHE.get(line_key, []))
        if snapshot:
            yield f"data: {json.dumps(snapshot)}\n\n"

        try:
            while True:
                try:
                    data = client_q.get(timeout=30)
                    yield f"data: {json.dumps(data)}\n\n"
                except queue.Empty:
                    yield ": keepalive\n\n"  # Prevent proxy timeout
        except GeneratorExit:
            pass
        finally:
            with SSE_LOCK:
                if client in SSE_CLIENTS:
                    SSE_CLIENTS.remove(client)

    return Response(
        stream(),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


# =========================================================
# BACKGROUND SERVICES
# =========================================================
def _shutdown_polling():
    POLLING.stop(timeout=3)
    POLL_LEASE.release()


def start_background_services():
    """Start the fixture worker and the managed Modbus polling.

    Polling needs the single-owner lease so that two processes on this PC can
    never poll the same USRs. start() is idempotent, so calling this again
    cannot create duplicate workers.
    """
    threading.Thread(target=fixture_worker, daemon=True, name="fixture-worker").start()

    if not POLL_LEASE.acquire():
        logger.error(
            "Another process on this PC already holds the polling lease (port %s) — "
            "this process will NOT poll the gauges. Stop the other instance first.",
            SETTINGS.lease_port,
        )
        return
    started = POLLING.start()
    atexit.register(_shutdown_polling)
    logger.info(
        "Polling started: mode=%s lines=%s workers=%s",
        SETTINGS.polling_mode,
        ",".join(sorted(SETTINGS.poll_lines)),
        len(started),
    )


def _is_reloader_parent():
    """Flask's debug reloader imports this module twice; only the child serves requests."""
    return (
        os.environ.get("FLASK_DEBUG", "false").lower() == "true"
        and os.environ.get("WERKZEUG_RUN_MAIN") != "true"
    )


if "pytest" not in sys.modules:
    for _warning in warnings_for(CONFIG):
        logger.warning(_warning)
    logger.info(
        "Lines loaded from .env: %s",
        ", ".join(
            f"{ln.key} ({len(ln.gauge_ids)} gauges, {len(ln.gateways)} USRs)"
            for ln in LINES.values()
        ),
    )
    if not _is_reloader_parent():
        try:
            ensure_test_header_columns()
        except Exception:
            logger.critical("Database unavailable at startup", exc_info=True)

if (
    "pytest" not in sys.modules
    and os.environ.get("DISABLE_WORKERS", "0") != "1"
    and not _is_reloader_parent()
):
    start_background_services()


# =========================================================
# LIVE GAUGE STATUS API
# =========================================================
@app.route("/api/fixture-live/<int:slave_id>")
@login_required
def get_live_vacuum(slave_id):
    snap = POLLING.gauge_snapshot(g.line.key, slave_id)
    if snap is None:
        return jsonify({"error": "Unknown gauge"}), 404
    # `vacuum` is only a trustworthy current reading (None when stale or
    # unavailable); `last_value` is the last known value for greyed-out display.
    snap["vacuum"] = snap["value"]
    return jsonify(snap)


# =========================================================
# FIXTURE DETAIL API (extended)
# =========================================================
@app.route("/api/fixture/<int:slave_id>")
@login_required
@db_safe
def fixture_detail(slave_id):
    conn = get_connection()
    cur = conn.cursor()
    cur.execute(
        """
        SELECT TOP 1
            h.gauge_id,
            h.serial_no,
            h.model_code,
            h.model_name,
            h.line_name,
            h.final_result,
            h.start_time,
            h.upper_limit
        FROM pirani_test_header h
        WHERE h.gauge_id = ? AND h.line_name = ?
        ORDER BY h.start_time DESC
    """,
        (slave_id, g.line.key),
    )
    row = cur.fetchone()
    conn.close()

    if not row:
        return jsonify(
            {
                "gauge_id": slave_id,
                "serial_no": None,
                "model_code": None,
                "model_name": None,
                "line_name": None,
                "final_result": None,
                "start_time": None,
                "ul": None,
            }
        )

    return jsonify(
        {
            "gauge_id": row[0],
            "serial_no": row[1],
            "model_code": row[2],
            "model_name": row[3],
            "line_name": row[4],
            "final_result": row[5],
            "start_time": str(row[6]) if row[6] else None,
            "ul": float(row[7]) if row[7] is not None else None,
        }
    )


# =========================================================
# MODBUS DIAGNOSTICS / WORKER HEALTH (own line only)
# =========================================================
@app.route("/api/modbus/diagnostics")
@login_required
def modbus_diagnostics():
    gateways = []
    for usr in POLLING.line_snapshot(g.line.key):
        gauges = usr["gauges"]
        usr["gauges_polled"] = len(gauges)
        usr["gauges_responding"] = sum(1 for gs in gauges.values() if gs["freshness"] == "valid")
        usr["readings"] = {
            sid: round(gs["value"], 3) if gs["value"] is not None else None
            for sid, gs in gauges.items()
        }
        gateways.append(usr)
    return jsonify({"gateways": gateways, "polling_mode": SETTINGS.polling_mode})


# =========================================================
# LINE CONFIGURATION (read-only — it lives in .env)
# =========================================================
def _line_settings_payload(line):
    return {
        "line_key": line.key,
        "line_label": line.label,
        "gap": line.gap,
        "reading_delay_sec": line.reading_delay_sec,
        "upper_limit": line.upper_limit,
        "poll_interval_sec": line.poll_interval_sec,
        "min_duration_sec": line.min_duration_sec,
        "register_map": line.register_map,
        "polling_mode": SETTINGS.polling_mode,
        "polled_here": POLLING.polled_here(line.key),
        "network": {
            "subnet": line.subnet,
            "pc_ip": line.pc_line_ip,
            "router_ip": line.router_ip,
            "company_pc_ip": line.company_pc_ip,
        },
        "hosts": [
            {
                "usr_id": gw.usr_id,
                "enabled": gw.enabled,
                "host": gw.host,
                "port": gw.port,
                "first_gauge": gw.first,
                "last_gauge": gw.last,
                "gauges": [
                    {
                        "id": n,
                        "name": f"Gauge {n}",
                        "slave_id": n,
                        "enabled": n not in line.disabled_gauges,
                    }
                    for n in gw.gauge_numbers
                ],
            }
            for gw in line.gateways
        ],
    }


@app.route("/api/settings/lines", methods=["GET"])
@login_required
def settings_lines():
    # Only the logged-in line — never another line's USRs or settings.
    return jsonify([_line_settings_payload(g.line)])


def _host_has_active_test(line, gateway):
    """True if any gauge behind this gateway has a test running right now.
    Guards test-connection from stealing the TCP connection out from under a
    real test on gateways that only support one connection at a time (common
    with cheap Modbus TCP-to-RTU gateways)."""
    return bool(set(gateway.gauge_numbers) & set(get_active_tests(line.key)))


def _test_modbus_connection(host, port, slave_ids, reg_map, timeout=3):
    """One-shot connectivity check, independent of the polling workers.
    Connects, then (if any slave_ids given) reads each gauge to confirm the
    gauges themselves respond, not just the gateway's TCP port."""
    client = ModbusTcpClient(host=host, port=port, timeout=timeout)
    try:
        if not client.connect():
            return {"reachable": False, "error": "Connection timed out or refused", "gauges": []}

        gauges = []
        for sid in slave_ids:
            r = read_gauge(client, sid, reg_map)
            gauges.append(
                {
                    "slave_id": sid,
                    "responding": r.status == SUCCESS,
                    "value": round(r.value, 3) if r.value is not None else None,
                    "communication_status": r.status,
                }
            )
        return {"reachable": True, "error": None, "gauges": gauges}
    except Exception as e:
        return {"reachable": False, "error": str(e), "gauges": []}
    finally:
        try:
            client.close()
        except Exception:
            pass


@app.route("/api/settings/test-connection", methods=["POST"])
@login_required
@admin_required
def test_connection_route():
    line = g.line
    d = request.get_json(silent=True) or {}
    host = (d.get("host") or "").strip()
    port = d.get("port")
    slave_ids = d.get("slave_ids") or []

    if not host or not isinstance(port, int) or isinstance(port, bool) or not (1 <= port <= 65535):
        return jsonify({"error": "Invalid host/port"}), 400
    if not isinstance(slave_ids, list) or not all(
        isinstance(s, int) and not isinstance(s, bool) for s in slave_ids
    ):
        return jsonify({"error": "slave_ids must be a list of integers"}), 400

    # Only this line's own USRs may be probed — this endpoint must not be
    # usable to connect to arbitrary hosts.
    gateway = next((gw for gw in line.gateways if gw.host == host and gw.port == port), None)
    if gateway is None:
        return jsonify({"error": "Not a gateway of your line"}), 403
    if any(s not in gateway.gauge_numbers for s in slave_ids):
        return jsonify({"error": "slave_ids must belong to this gateway"}), 400

    if not POLLING.polled_here(line.key):
        return jsonify({
            "error": f"{line.label} is not reachable from this PC "
                     f"(polling mode: {SETTINGS.polling_mode}) — run the test on the line PC",
        }), 409

    if _host_has_active_test(line, gateway):
        return jsonify({
            "error": "Cannot test — an active test is currently using this gateway",
        }), 409

    return jsonify(_test_modbus_connection(host, port, slave_ids, REGISTER_MAPS[line.register_map]))


# =========================================================
# REPORTS API
# =========================================================
@app.route("/api/reports")
@login_required
@db_safe
def reports_api():
    line = g.line

    # The line is part of the cache key: two lines can send identical query
    # strings, and must never be served each other's cached rows.
    cache_key = (line.key,) + tuple(sorted(request.args.items()))
    if cache_key in REPORT_CACHE and request.args.get("export") != "excel":
        return jsonify(REPORT_CACHE[cache_key])

    limit = request.args.get("limit", type=int)
    start = request.args.get("start")
    end = request.args.get("end")
    model = request.args.get("model")
    result = request.args.get("result")
    gauge = request.args.get("gauge", type=int)

    base_query = """
        SELECT
            h.test_id,
            h.gauge_id,
            h.serial_no,
            h.model_code,
            h.model_name,
            h.line_name,
            h.final_result,
            h.start_time,
            h.end_time,
            l.vacuum AS last_vacuum,
            h.upper_limit AS ul
        FROM pirani_test_header h
        OUTER APPLY (
            SELECT TOP 1 vacuum
            FROM pirani_test_log
            WHERE test_id = h.test_id
            ORDER BY log_time DESC
        ) l
    """
    # Always restricted to the logged-in line; a `line` query parameter from
    # the client is ignored.
    params = [line.key]

    if limit:
        query = (
            base_query + " WHERE h.line_name = ? ORDER BY h.start_time DESC "
            "OFFSET 0 ROWS FETCH NEXT ? ROWS ONLY"
        )
        params.append(limit)
    else:
        if not start or not end:
            return jsonify([])

        # Use CAST so the full end-date is included
        query = (
            base_query
            + " WHERE h.line_name = ? AND CAST(h.start_time AS DATE) BETWEEN ? AND ?"
        )
        params.extend([start, end])

        if model:
            query += " AND h.model_code = ?"
            params.append(model)
        if result:
            query += " AND h.final_result = ?"
            params.append(result)
        if gauge:
            query += " AND h.gauge_id = ?"
            params.append(gauge)

        query += " ORDER BY h.start_time DESC"

    conn = get_connection()
    df = pd.read_sql(query, conn, params=params)
    conn.close()

    if "test_id" in df.columns:
        df["test_id"] = df["test_id"].astype(str)

    data = df.to_dict(orient="records")

    # A running test has no end_time yet (and often no final_result / last
    # reading). Depending on pandas dtype inference these missing values can
    # come through as NaN/NaT rather than None, and json.dumps emits those
    # as a bare `NaN` token, which is NOT valid JSON and breaks the browser's
    # JSON.parse(). pd.isna() reliably detects NaN/NaT/None/pd.NA regardless
    # of column dtype, so normalize everything to a real `None` here.
    for row in data:
        for key, value in row.items():
            if pd.isna(value):
                row[key] = None
            elif isinstance(value, pd.Timestamp):
                row[key] = str(value)

    REPORT_CACHE[cache_key] = data

    if request.args.get("export") == "excel":
        out = io.BytesIO()
        # Reserve the top rows for the logo; the real header/data starts below it.
        df.to_excel(out, index=False, startrow=3)
        out.seek(0)

        if os.path.exists(LOGO_PATH):
            wb = openpyxl.load_workbook(out)
            ws = wb.active
            img = XLImage(LOGO_PATH)
            img.width, img.height = 120, 70
            ws.add_image(img, "A1")
            out = io.BytesIO()
            wb.save(out)
            out.seek(0)

        fname = f"pirani_report_{start or 'all'}.xlsx"
        return send_file(
            out,
            as_attachment=True,
            download_name=fname,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    return jsonify(data)


# =========================================================
# REPORT TREND API  (UUID-safe — was incorrectly int:)
# =========================================================
@app.route("/api/report/<string:test_id>/trend")
@login_required
@db_safe
def report_trend(test_id):
    conn = get_connection()
    cur = conn.cursor()
    # Joined to the header so a test of another line reads as "no readings".
    cur.execute(
        """
        SELECT l.log_time, l.vacuum, l.result
        FROM pirani_test_log l
        JOIN pirani_test_header h ON h.test_id = l.test_id
        WHERE l.test_id = ? AND h.line_name = ?
        ORDER BY l.log_time
    """,
        (test_id, g.line.key),
    )
    rows = cur.fetchall()
    conn.close()

    return jsonify(
        [
            {
                "time": str(r[0]),
                "vacuum": float(r[1]) if r[1] is not None else None,
                "result": r[2],
            }
            for r in rows
        ]
    )


# =========================================================
# REPORT PDF EXPORT  (single test: header info + every poll reading)
# =========================================================
def _trend_chart_drawing(readings, ul, width=480, height=190):
    """Vacuum Trend Chart
    X-axis : Time (min)
    Y-axis : Vacuum (mbar)
    Blue   : Actual Vacuum
    Red    : Upper Limit
    """

    if len(readings) < 2:
        return None

    # -----------------------------
    # Convert timestamps to elapsed minutes
    # -----------------------------
    try:
        start = datetime.fromisoformat(str(readings[0][0]))
    except Exception:
        start = None

    x_labels = []
    values = []

    for r in readings:
        if r[1] is None:
            continue

        values.append(float(r[1]))

        if start:
            try:
                t = datetime.fromisoformat(str(r[0]))
                mins = (t - start).total_seconds() / 60.0
                x_labels.append(f"{mins:.1f}")
            except Exception:
                x_labels.append(str(len(x_labels)))
        else:
            x_labels.append(str(len(x_labels)))

    n = len(values)

    if n < 2:
        return None

    # -----------------------------
    # Axis limits
    # -----------------------------
    bounds = values.copy()

    if ul is not None:
        bounds.append(float(ul))

    pad = (max(bounds) - min(bounds)) * 0.15 or 0.05

    y_min = min(bounds) - pad
    y_max = max(bounds) + pad

    drawing = Drawing(width, height)

    chart = HorizontalLineChart()
    chart.x = 55
    chart.y = 40
    chart.width = width - 80
    chart.height = height - 70

    chart.valueAxis.valueMin = y_min
    chart.valueAxis.valueMax = y_max
    chart.valueAxis.labelTextFormat = "%.3f"

    chart.valueAxis.labels.fontSize = 7
    chart.categoryAxis.labels.fontSize = 7

    # Show only every few labels
    step = max(1, n // 8)

    chart.categoryAxis.categoryNames = [
        x_labels[i] if i % step == 0 else "" for i in range(n)
    ]

    # -----------------------------
    # Data series
    # -----------------------------
    chart.data = [
        values,  # Blue
        [float(ul)] * n if ul is not None else [],
    ]

    # Vacuum
    chart.lines[0].strokeColor = colors.blue
    chart.lines[0].strokeWidth = 2
    chart.lines[0].symbol = None

    # Upper Limit
    if ul is not None:
        chart.lines[1].strokeColor = colors.red
        chart.lines[1].strokeWidth = 1.5
        chart.lines[1].strokeDashArray = (4, 2)
        chart.lines[1].symbol = None

    drawing.add(chart)

    # -----------------------------
    # Axis Titles
    # -----------------------------
    drawing.add(
        String(
            width / 2,
            8,
            "Time (min)",
            textAnchor="middle",
            fontSize=8,
        )
    )

    drawing.add(
        String(
            12,
            height / 2,
            "Vacuum (mbar)",
            angle=90,
            fontSize=8,
        )
    )

    # -----------------------------
    # Legend
    # -----------------------------
    legend = Legend()
    legend.x = width - 90
    legend.y = height - 10
    legend.dx = 10
    legend.dy = 10
    legend.fontSize = 7

    legend.colorNamePairs = [
        (colors.blue, "Vacuum"),
        (colors.red, "Upper Limit"),
    ]

    drawing.add(legend)

    return drawing


@app.route("/api/report/<string:test_id>/pdf")
@login_required
@db_safe
def report_pdf(test_id):
    conn = get_connection()
    cur = conn.cursor()

    cur.execute(
        """
        SELECT h.test_id, h.gauge_id, h.serial_no, h.model_code, h.model_name,
               h.line_name, h.final_result, h.start_time, h.end_time,
               h.upper_limit AS ul
        FROM pirani_test_header h
        WHERE h.test_id = ? AND h.line_name = ?
    """,
        (test_id, g.line.key),
    )
    header_row = cur.fetchone()
    if not header_row:
        conn.close()
        return jsonify({"error": "Test not found"}), 404
    header = dict(zip([c[0] for c in cur.description], header_row))

    cur.execute(
        """
        SELECT log_time, vacuum, result
        FROM pirani_test_log
        WHERE test_id = ?
        ORDER BY log_time
    """,
        (test_id,),
    )
    readings = cur.fetchall()
    conn.close()

    def fmt_dt(v):
        return str(v) if v else "—"

    def fmt_num(v):
        return f"{float(v):.3f}" if v is not None else "—"

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        topMargin=18 * mm,
        bottomMargin=16 * mm,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Title"], fontSize=16, spaceAfter=4
    )
    sub_style = ParagraphStyle(
        "ReportSub",
        parent=styles["Normal"],
        textColor=colors.HexColor("#64748b"),
        fontSize=9,
    )
    h2_style = ParagraphStyle(
        "ReportH2", parent=styles["Heading2"], fontSize=12, spaceBefore=2, spaceAfter=6
    )

    elements = []
    if os.path.exists(LOGO_PATH):
        logo = Image(LOGO_PATH, width=32 * mm, height=18.8 * mm)
        elements += [logo, Spacer(1, 6)]
    elements += [
        Paragraph("Pirani Gauge Test Report", title_style),
        Paragraph(f"Test ID: {header['test_id']}", sub_style),
        Spacer(1, 10),
    ]

    info_rows = [
        [
            "Gauge",
            str(header.get("gauge_id") or "—"),
            "Line",
            header.get("line_name") or "—",
        ],
        [
            "Serial No.",
            header.get("serial_no") or "—",
            "Model",
            header.get("model_code") or "—",
        ],
        [
            "Model Name",
            header.get("model_name") or "—",
            "Result",
            header.get("final_result") or "—",
        ],
        [
            "Start Time",
            fmt_dt(header.get("start_time")),
            "End Time",
            fmt_dt(header.get("end_time")),
        ],
        [
            "Upper Limit",
            fmt_num(header.get("ul")),
            "",
            "",
        ],
    ]
    info_table = Table(info_rows, colWidths=[28 * mm, 55 * mm, 28 * mm, 55 * mm])
    info_table.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#64748b")),
                ("TEXTCOLOR", (2, 0), (2, -1), colors.HexColor("#64748b")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f8fafc")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#f8fafc")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f0")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    elements += [info_table, Spacer(1, 14)]

    vacs = [float(r[1]) for r in readings if r[1] is not None]
    if vacs:
        stats_table = Table(
            [
                ["Min (mbar)", "Max (mbar)", "Avg (mbar)", "Readings"],
                [
                    f"{min(vacs):.3f}",
                    f"{max(vacs):.3f}",
                    f"{(sum(vacs) / len(vacs)):.3f}",
                    str(len(vacs)),
                ],
            ],
            colWidths=[41.5 * mm] * 4,
        )
        stats_table.setStyle(
            TableStyle(
                [
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f0")),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        elements += [stats_table, Spacer(1, 14)]

    chart_drawing = _trend_chart_drawing(readings, header.get("ul"))
    if chart_drawing:
        elements.append(Paragraph("Vacuum Trend", h2_style))
        elements.append(chart_drawing)
        elements.append(Spacer(1, 14))

    elements.append(Paragraph("Readings Log", h2_style))

    table_data = [["#", "Time", "Vacuum (mbar)", "Upper Limit (mbar)", "Result"]]

    for i, r in enumerate(readings, start=1):
        table_data.append(
            [
                str(i),
                str(r[0]),
                f"{float(r[1]):.3f}" if r[1] is not None else "—",
                fmt_num(header.get("ul")),
                r[2] or "—",
            ]
        )

    if len(table_data) == 1:
        table_data.append(["—", "No readings recorded", "", "", ""])

    readings_table = Table(
        table_data,
        colWidths=[
            10 * mm,  # #
            50 * mm,  # Time
            33 * mm,  # Vacuum
            30 * mm,  # Upper Limit
            22 * mm,  # Result
        ],
        repeatRows=1,
    )

    readings_table.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#e2e8f0")),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#f8fafc")],
                ),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    elements.append(readings_table)

    doc.build(elements)
    buf.seek(0)

    return send_file(
        buf,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"pirani_report_{test_id}.pdf",
    )


# =========================================================
# APP START
# =========================================================
if __name__ == "__main__":
    debug = os.environ.get("FLASK_DEBUG", "false").lower() == "true"
    app.run(host="0.0.0.0", port=5000, threaded=True, debug=debug)
